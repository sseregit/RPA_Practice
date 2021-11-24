from openpyxl import load_workbook

wb = load_workbook("xlsx/sample.xlsx")
ws = wb.active

# 차트를 그리기위해 import
from openpyxl.chart import BarChart, Reference, LineChart
# 어떤 데이터를 차트를 만들지 정해야함.
# B2:C11
bar_value = Reference(ws,min_row=2, max_row=11, min_col=2, max_col=3)
# chart의 종류
bar_chart = BarChart()
# chart data 추가
bar_chart.add_data(bar_value)
# 차트 넣을 위치 정의
ws.add_chart(bar_chart, "E1")

# B1:C11 까지 데이터
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
# titles_from_data 계열 > 제목이 정해진다.
line_chart.add_data(line_value, titles_from_data=True)
# 제목
line_chart.title = "성적표"
# 미리 정의된 스타일 적용
line_chart.style = 10
# Y축의 제목
line_chart.y_axis.title = "점수"
# X축의 제목
line_chart.x_axis.title = "번호"
ws.add_chart(line_chart, "N1")

wb.save("xlsx/sample_barchart.xlsx")